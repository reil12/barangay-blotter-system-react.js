from flask import (
    Flask, request, session, make_response, jsonify,
    send_from_directory, abort
)
import pyodbc
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
from flask_mail import Mail, Message
import random
import string
import os
from datetime import datetime, date, timedelta
import io
from html import escape

app = Flask(__name__)
app.secret_key = 'barangay_secret_key_2026'

# The React app is built into static/react (see README_REACT.md)
REACT_DIR = os.path.join(app.static_folder, 'react')

# =========================================================
# MAIL CONFIGURATION (GMAIL)
# =========================================================
app.config['MAIL_SERVER'] = 'smtp.gmail.com'
app.config['MAIL_PORT'] = 587
app.config['MAIL_USE_TLS'] = True
app.config['MAIL_USERNAME'] = 'reilestrada8@gmail.com'
app.config['MAIL_PASSWORD'] = 'rzjiqqzfazzdrrrq'
app.config['MAIL_DEFAULT_SENDER'] = 'reilestrada8@gmail.com'
app.config['MAIL_TIMEOUT'] = 10

mail = Mail(app)

# =========================================================
# DATABASE CONNECTION
# =========================================================
def get_db():
    return pyodbc.connect(
        "DRIVER={ODBC Driver 17 for SQL Server};"
        "SERVER=ARJING\\SQLEXPRESS;"
        "DATABASE=BarangayBlotterDB;"
        "Trusted_Connection=yes;"
    )

# =========================================================
# HELPERS
# =========================================================
def generate_otp():
    return ''.join(random.choices(string.digits, k=6))

def generate_case_number():
    conn = get_db()
    cursor = conn.cursor()
    year = datetime.now().year
    cursor.execute("""
        SELECT COUNT(*) FROM Incidents
        WHERE YEAR(DateFiled) = ?
    """, year)
    count = cursor.fetchone()[0] + 1
    conn.close()
    return f"BRG-{year}-{str(count).zfill(4)}"

def rows_to_dicts(cursor):
    """Turn every row of the last query into a JSON-friendly dict."""
    cols = [c[0] for c in cursor.description]
    out = []
    for row in cursor.fetchall():
        d = dict(zip(cols, row))
        for k, v in d.items():
            if isinstance(v, (datetime, date)):
                d[k] = v.isoformat()
        out.append(d)
    return out

def one_to_dict(cursor):
    rows = rows_to_dicts(cursor)
    return rows[0] if rows else None

def otp_email_html(full_name, otp, intro, expiry_text):
    return f"""
    <div style="font-family:Arial;padding:20px;
                max-width:480px;margin:auto;
                border:1px solid #ddd;border-radius:10px;">
        <div style="background:#1a3c5e;padding:20px;
                    border-radius:8px 8px 0 0;
                    text-align:center;">
            <h2 style="color:white;margin:0;">
                Barangay Blotter System
            </h2>
        </div>
        <div style="padding:24px;">
            <p>Hello <strong>{full_name}</strong>,</p>
            <p>{intro}</p>
            <div style="font-size:38px;font-weight:bold;
                        letter-spacing:12px;color:#1a3c5e;
                        background:#f0f4f8;padding:20px;
                        text-align:center;border-radius:8px;
                        margin:20px 0;">{otp}</div>
            <p style="color:#888;font-size:13px;">
                {expiry_text}
            </p>
        </div>
    </div>
    """

def meeting_email_html(complainant_name, case_number, respondent, meet_date, meet_time, place, notes):
    notes_block = ''
    if notes:
        notes_block = f"<p><strong>Additional instructions:</strong><br>{escape(notes)}</p>"
    return f"""
    <div style="font-family:Arial,Helvetica,sans-serif;max-width:560px;margin:auto;color:#1f2937;">
        <div style="background:#1f3a5f;color:#ffffff;padding:16px 22px;">
            <strong>Barangay Blotter System</strong><br>
            <span style="font-size:13px;">Notice of Mediation Meeting</span>
        </div>
        <div style="border:1px solid #d1d5db;border-top:none;padding:22px;">
            <p>Dear <strong>{escape(complainant_name)}</strong>,</p>
            <p>This is to notify you of a barangay mediation meeting regarding
               Case No. <strong>{escape(case_number)}</strong>, in which
               <strong>{escape(respondent)}</strong> is the reported person.</p>
            <table style="border-collapse:collapse;margin:14px 0;width:100%;">
                <tr><td style="padding:8px;border:1px solid #d1d5db;width:110px;background:#f3f4f6;"><strong>Date</strong></td>
                    <td style="padding:8px;border:1px solid #d1d5db;">{escape(meet_date)}</td></tr>
                <tr><td style="padding:8px;border:1px solid #d1d5db;background:#f3f4f6;"><strong>Time</strong></td>
                    <td style="padding:8px;border:1px solid #d1d5db;">{escape(meet_time)}</td></tr>
                <tr><td style="padding:8px;border:1px solid #d1d5db;background:#f3f4f6;"><strong>Place</strong></td>
                    <td style="padding:8px;border:1px solid #d1d5db;">{escape(place)}</td></tr>
            </table>
            {notes_block}
            <p>Please attend on time and bring a valid ID and any supporting documents.</p>
            <p style="color:#6b7280;font-size:13px;margin-top:24px;">
                This is an automated message from the Barangay Blotter System. For questions, please visit the Barangay Hall during office hours.
            </p>
        </div>
    </div>
    """

# =========================================================
# DECORATORS (JSON API)
# =========================================================
def api_login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return jsonify(error='Not logged in'), 401
        return f(*args, **kwargs)
    return decorated

def api_admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return jsonify(error='Not logged in'), 401
        if session.get('role') != 'admin':
            return jsonify(error='Forbidden'), 403
        return f(*args, **kwargs)
    return decorated

# =========================================================
# AUTH
# =========================================================
@app.route('/api/me')
def api_me():
    """Who is logged in? Also tells React if an OTP step is pending."""
    user = None
    if 'user_id' in session:
        user = {
            'id': session['user_id'],
            'name': session.get('user_name'),
            'role': session.get('role'),
        }
    pending = None
    if 'login_otp' in session:
        pending = 'login'
    elif 'verify_email' in session:
        pending = 'signup'
    return jsonify(user=user, pending=pending)

@app.route('/api/login', methods=['POST'])
def api_login():
    body = request.get_json(silent=True) or {}
    email = body.get('email', '')
    password = body.get('password', '')
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM Users WHERE Email=?", email)
    user = cursor.fetchone()
    conn.close()
    if not (user and check_password_hash(user.PasswordHash, password)):
        return jsonify(error='Invalid email or password.'), 401
    if user.IsVerified == 0 or user.IsVerified is None:
        return jsonify(error='Please verify your email before logging in.'), 403
    otp = generate_otp()
    session['login_otp'] = otp
    session['login_otp_expiry'] = (
        datetime.now() + timedelta(minutes=10)
    ).strftime('%Y-%m-%d %H:%M:%S')
    session['pending_user_id'] = user.UserID
    session['pending_user_name'] = user.FullName
    session['pending_role'] = user.Role
    try:
        msg = Message('Your Barangay Blotter Login Code', recipients=[email])
        msg.html = otp_email_html(
            user.FullName, otp,
            'Your login verification code is:',
            'This code expires in <strong>10 minutes</strong>.'
        )
        mail.send(msg)
    except Exception as e:
        return jsonify(error=f'Could not send email: {str(e)}'), 500
    return jsonify(ok=True, message='A 6-digit code was sent to your email.')

@app.route('/api/verify-login', methods=['POST'])
def api_verify_login():
    if 'login_otp' not in session:
        return jsonify(error='No login in progress.', redirect='login'), 400
    entered = (request.get_json(silent=True) or {}).get('otp', '')
    expiry = datetime.strptime(
        session['login_otp_expiry'], '%Y-%m-%d %H:%M:%S'
    )
    if datetime.now() > expiry:
        for k in ['login_otp', 'login_otp_expiry',
                  'pending_user_id', 'pending_user_name', 'pending_role']:
            session.pop(k, None)
        return jsonify(
            error='Code expired. Please log in again.', redirect='login'
        ), 400
    if entered == session['login_otp']:
        session['user_id'] = session.pop('pending_user_id')
        session['user_name'] = session.pop('pending_user_name')
        session['role'] = session.pop('pending_role')
        session.pop('login_otp', None)
        session.pop('login_otp_expiry', None)
        return jsonify(ok=True, user={
            'id': session['user_id'],
            'name': session['user_name'],
            'role': session['role'],
        })
    return jsonify(error='Incorrect verification code.'), 400

@app.route('/api/signup', methods=['POST'])
def api_signup():
    data = request.get_json(silent=True) or {}
    fullname = data.get('fullname', '')
    email = data.get('email', '')
    password = data.get('password', '')
    address = data.get('address', '')
    contact = data.get('contact', '')
    if not (fullname and email and password):
        return jsonify(error='Please fill in all required fields.'), 400
    hashed_password = generate_password_hash(password)
    otp = generate_otp()
    otp_expiry = (
        datetime.now() + timedelta(minutes=10)
    ).strftime('%Y-%m-%d %H:%M:%S')
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM Users WHERE Email=?", email)
    if cursor.fetchone():
        conn.close()
        return jsonify(error='Email already registered.'), 409
    try:
        msg = Message('Verify Your Barangay Blotter Account', recipients=[email])
        msg.html = otp_email_html(
            fullname, otp,
            'Your verification code is:',
            'Expires in <strong>10 minutes</strong>.'
        )
        mail.send(msg)
        cursor.execute("""
            INSERT INTO Users
            (FullName, Email, PasswordHash, Address,
             ContactNumber, Role, IsVerified, VerifyOTP, OTPExpiry)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, fullname, email, hashed_password, address,
             contact, 'user', 0, otp, otp_expiry)
        conn.commit()
        conn.close()
        session['verify_email'] = email
        return jsonify(ok=True, message='Account created! Check your email.')
    except Exception as e:
        conn.close()
        return jsonify(error=f'Error: {str(e)}'), 500

@app.route('/api/verify-signup', methods=['POST'])
def api_verify_signup():
    if 'verify_email' not in session:
        return jsonify(error='No signup in progress.', redirect='signup'), 400
    entered_otp = (request.get_json(silent=True) or {}).get('otp', '')
    email = session['verify_email']
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT VerifyOTP, OTPExpiry FROM Users WHERE Email=?", email
    )
    row = cursor.fetchone()
    if not row:
        conn.close()
        return jsonify(error='Account not found.', redirect='signup'), 400
    expiry = datetime.strptime(str(row.OTPExpiry)[:19], '%Y-%m-%d %H:%M:%S')
    if datetime.now() > expiry:
        cursor.execute("DELETE FROM Users WHERE Email=?", email)
        conn.commit()
        conn.close()
        session.pop('verify_email', None)
        return jsonify(
            error='OTP expired. Please register again.', redirect='signup'
        ), 400
    if entered_otp == row.VerifyOTP:
        cursor.execute("UPDATE Users SET IsVerified=1 WHERE Email=?", email)
        conn.commit()
        conn.close()
        session.pop('verify_email', None)
        return jsonify(ok=True, message='Email verified! You can now log in.')
    conn.close()
    return jsonify(error='Incorrect OTP code.'), 400

@app.route('/api/logout', methods=['POST'])
def api_logout():
    session.clear()
    return jsonify(ok=True)

# =========================================================
# RESIDENT (USER) API
# =========================================================
@app.route('/api/my-reports')
@api_login_required
def api_my_reports():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT IncidentID, CaseNumber, Complainant, Respondent,
               IncidentType, DateFiled, Status
        FROM Incidents WHERE UserID=? ORDER BY DateFiled DESC
    """, session['user_id'])
    data = rows_to_dicts(cursor)
    conn.close()
    return jsonify(data)

@app.route('/api/case-number')
@api_login_required
def api_case_number():
    return jsonify(caseNumber=generate_case_number())

@app.route('/api/incidents', methods=['POST'])
@api_login_required
def api_submit_incident():
    data = request.get_json(silent=True) or {}
    required = ['complainant', 'respondent', 'incident_type',
                'date_incident', 'narrative']
    if not all(data.get(k) for k in required):
        return jsonify(error='Please fill in all required fields.'), 400
    # Use the case number shown on the form; generate one if it is missing
    case_number = data.get('case_number') or generate_case_number()
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO Incidents
        (CaseNumber, Complainant, Respondent,
         IncidentType, DateIncident, Narrative,
         Status, UserID)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, case_number, data['complainant'],
         data['respondent'], data['incident_type'],
         data['date_incident'], data['narrative'],
         'Pending', session['user_id'])
    conn.commit()
    conn.close()
    return jsonify(ok=True, caseNumber=case_number,
                   message='Blotter report submitted successfully.')

@app.route('/api/announcements')
@api_login_required
def api_announcements():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM Announcements ORDER BY CreatedAt DESC")
    data = rows_to_dicts(cursor)
    conn.close()
    return jsonify(data)

# =========================================================
# ADMIN API: INCIDENTS
# =========================================================
@app.route('/api/admin/incidents')
@api_admin_required
def api_admin_incidents():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT IncidentID, CaseNumber, Complainant, Respondent,
               IncidentType, DateFiled, Status
        FROM Incidents ORDER BY DateFiled DESC
    """)
    data = rows_to_dicts(cursor)
    conn.close()
    return jsonify(data)

@app.route('/api/admin/incidents/<int:id>', methods=['GET'])
@api_admin_required
def api_admin_incident(id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM Incidents WHERE IncidentID=?", id)
    record = one_to_dict(cursor)
    conn.close()
    if not record:
        return jsonify(error='Record not found'), 404
    return jsonify(record)

@app.route('/api/admin/incidents/<int:id>', methods=['PUT'])
@api_admin_required
def api_admin_edit_incident(id):
    data = request.get_json(silent=True) or {}
    if data.get('status') not in ('Pending', 'Settled', 'Dismissed'):
        return jsonify(error='Invalid status'), 400
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        UPDATE Incidents
        SET Complainant=?, Respondent=?,
            IncidentType=?, DateIncident=?,
            Narrative=?, Status=?
        WHERE IncidentID=?
    """, data.get('complainant'), data.get('respondent'),
         data.get('incident_type'), data.get('date_incident') or None,
         data.get('narrative'), data['status'], id)
    conn.commit()
    conn.close()
    return jsonify(ok=True, message='Blotter record updated successfully.')

@app.route('/api/admin/incidents/<int:id>/status', methods=['POST'])
@api_admin_required
def api_update_status(id):
    status = (request.get_json(silent=True) or {}).get('status')
    if status not in ('Pending', 'Settled', 'Dismissed'):
        return jsonify(error='Invalid status'), 400
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("UPDATE Incidents SET Status=? WHERE IncidentID=?", status, id)
    conn.commit()
    conn.close()
    return jsonify(ok=True, status=status)

@app.route('/api/admin/incidents/<int:id>', methods=['DELETE'])
@api_admin_required
def api_delete_incident(id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM Incidents WHERE IncidentID=?", id)
    conn.commit()
    conn.close()
    return jsonify(ok=True)

@app.route('/api/admin/incidents/<int:id>/certificate')
@api_admin_required
def api_certificate(id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT i.*, u.FullName as FiledBy,
               u.Address as FiledByAddress,
               u.ContactNumber as FiledByContact
        FROM Incidents i
        LEFT JOIN Users u ON i.UserID = u.UserID
        WHERE i.IncidentID = ?
    """, id)
    record = one_to_dict(cursor)
    conn.close()
    if not record:
        return jsonify(error='Record not found'), 404
    return jsonify(record=record, now=datetime.now().isoformat())

# =========================================================
# ADMIN API: ANALYTICS
# =========================================================
@app.route('/api/admin/analytics')
@api_admin_required
def api_analytics():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT MONTH(DateFiled) as Month, COUNT(*) as Total
        FROM Incidents
        WHERE YEAR(DateFiled) = YEAR(GETDATE())
        GROUP BY MONTH(DateFiled)
        ORDER BY MONTH(DateFiled)
    """)
    monthly = [0] * 12
    for row in cursor.fetchall():
        monthly[row.Month - 1] = row.Total
    cursor.execute("""
        SELECT IncidentType, COUNT(*) as Total
        FROM Incidents
        GROUP BY IncidentType
        ORDER BY Total DESC
    """)
    types_raw = cursor.fetchall()
    incident_types = [row.IncidentType for row in types_raw]
    incident_counts = [row.Total for row in types_raw]

    def count(where=''):
        cursor.execute(f"SELECT COUNT(*) FROM Incidents {where}")
        return cursor.fetchone()[0]

    result = dict(
        monthly=monthly,
        incidentTypes=incident_types,
        incidentCounts=incident_counts,
        total=count(),
        pending=count("WHERE Status='Pending'"),
        settled=count("WHERE Status='Settled'"),
        dismissed=count("WHERE Status='Dismissed'"),
        thisWeek=count("WHERE DateFiled >= DATEADD(day, -7, GETDATE())"),
        lastWeek=count(
            "WHERE DateFiled >= DATEADD(day, -14, GETDATE()) "
            "AND DateFiled < DATEADD(day, -7, GETDATE())"
        ),
    )
    conn.close()
    return jsonify(result)

# =========================================================
# ADMIN: EXPORT EXCEL (file download)
# =========================================================
@app.route('/api/admin/export')
@api_admin_required
def export_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT i.IncidentID, i.CaseNumber, i.Complainant,
               i.Respondent, i.IncidentType, i.DateFiled,
               i.DateIncident, i.Narrative, i.Status,
               u.FullName as FiledBy
        FROM Incidents i
        LEFT JOIN Users u ON i.UserID = u.UserID
        ORDER BY i.DateFiled DESC
    """)
    records = cursor.fetchall()
    conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Blotter Records"
    header_fill = PatternFill("solid", fgColor="1a3c5e")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    headers = ['ID', 'Case No.', 'Complainant', 'Reported Person',
               'Type', 'Date Filed', 'Date of Incident',
               'Narrative', 'Status', 'Filed By']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    for row_idx, r in enumerate(records, 2):
        ws.cell(row=row_idx, column=1, value=r.IncidentID)
        ws.cell(row=row_idx, column=2, value=r.CaseNumber)
        ws.cell(row=row_idx, column=3, value=r.Complainant)
        ws.cell(row=row_idx, column=4, value=r.Respondent)
        ws.cell(row=row_idx, column=5, value=r.IncidentType)
        ws.cell(row=row_idx, column=6,
                value=r.DateFiled.strftime('%Y-%m-%d')
                if r.DateFiled else '')
        ws.cell(row=row_idx, column=7,
                value=str(r.DateIncident) if r.DateIncident else '')
        ws.cell(row=row_idx, column=8, value=r.Narrative)
        ws.cell(row=row_idx, column=9, value=r.Status)
        ws.cell(row=row_idx, column=10, value=r.FiledBy or 'N/A')
    col_widths = [6, 16, 20, 20, 20, 14, 16, 40, 12, 20]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[
            openpyxl.utils.get_column_letter(i)
        ].width = width
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    response = make_response(output.read())
    response.headers['Content-Disposition'] = (
        'attachment; filename=blotter_records.xlsx'
    )
    response.headers['Content-Type'] = (
        'application/vnd.openxmlformats-officedocument'
        '.spreadsheetml.sheet'
    )
    return response

# =========================================================
# ADMIN API: USERS
# =========================================================
@app.route('/api/admin/users')
@api_admin_required
def api_admin_users():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT u.UserID, u.FullName, u.Email,
               u.Address, u.ContactNumber,
               u.Role, u.IsVerified, u.CreatedAt,
               COUNT(i.IncidentID) AS TotalReports
        FROM Users u
        LEFT JOIN Incidents i ON u.UserID = i.UserID
        GROUP BY u.UserID, u.FullName, u.Email,
                 u.Address, u.ContactNumber,
                 u.Role, u.IsVerified, u.CreatedAt
        ORDER BY u.CreatedAt DESC
    """)
    data = rows_to_dicts(cursor)
    conn.close()
    return jsonify(data)

@app.route('/api/admin/users/<int:id>', methods=['DELETE'])
@api_admin_required
def api_delete_user(id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT Role FROM Users WHERE UserID=?", id)
    row = cursor.fetchone()
    if row and row.Role == 'admin':
        conn.close()
        return jsonify(error='Admin accounts are protected.'), 403
    cursor.execute("DELETE FROM Incidents WHERE UserID=?", id)
    cursor.execute("DELETE FROM Users WHERE UserID=?", id)
    conn.commit()
    conn.close()
    return jsonify(ok=True, message='User deleted successfully.')

# =========================================================
# ADMIN API: ANNOUNCEMENTS
# =========================================================
@app.route('/api/admin/announcements', methods=['POST'])
@api_admin_required
def api_add_announcement():
    data = request.get_json(silent=True) or {}
    title = data.get('title', '').strip()
    content = data.get('content', '').strip()
    category = data.get('category') or 'General'
    if not title or not content:
        return jsonify(error='Title and content are required.'), 400
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO Announcements (Title, Content, Category, AdminID)
        VALUES (?, ?, ?, ?)
    """, title, content, category, session['user_id'])
    conn.commit()
    conn.close()
    return jsonify(ok=True, message='Announcement posted successfully.')

@app.route('/api/admin/announcements/<int:id>', methods=['DELETE'])
@api_admin_required
def api_delete_announcement(id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM Announcements WHERE AnnouncementID=?", id)
    conn.commit()
    conn.close()
    return jsonify(ok=True, message='Announcement deleted.')

@app.route('/api/admin/incidents/<int:id>/meeting-notice', methods=['POST'])
@api_admin_required
def api_meeting_notice(id):
    """Email the complainant the date, time and place of a mediation meeting."""
    data = request.get_json(silent=True) or {}
    meet_date = (data.get('date') or '').strip()
    meet_time = (data.get('time') or '').strip()
    place = (data.get('place') or '').strip()
    notes = (data.get('notes') or '').strip()
    if not (meet_date and meet_time and place):
        return jsonify(error='Date, time and place are required.'), 400
    try:
        date_label = datetime.strptime(meet_date, '%Y-%m-%d').strftime('%B %d, %Y')
        time_label = datetime.strptime(meet_time, '%H:%M').strftime('%I:%M %p')
    except ValueError:
        return jsonify(error='Invalid date or time.'), 400

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT i.CaseNumber, i.Complainant, i.Respondent,
               u.Email, u.FullName
        FROM Incidents i
        LEFT JOIN Users u ON i.UserID = u.UserID
        WHERE i.IncidentID = ?
    """, id)
    rec = one_to_dict(cursor)
    conn.close()
    if not rec:
        return jsonify(error='Record not found'), 404
    if not rec['Email']:
        return jsonify(error='No email address on file for the complainant.'), 400

    try:
        msg = Message(
            f"Notice of Mediation Meeting - Case {rec['CaseNumber']}",
            recipients=[rec['Email']]
        )
        msg.html = meeting_email_html(
            rec['FullName'] or rec['Complainant'], rec['CaseNumber'],
            rec['Respondent'], date_label, time_label, place, notes
        )
        mail.send(msg)
    except Exception as e:
        return jsonify(error=f'Could not send email: {str(e)}'), 500
    return jsonify(ok=True, message=f"Meeting notice sent to {rec['Email']}.")

# =========================================================
# REACT APP (single-page app)
# Every non-API URL (/login, /dashboard, /admin/users ...) gets the
# same index.html; React Router decides which page to show.
# =========================================================
@app.route('/', defaults={'path': ''})
@app.route('/<path:path>')
def serve_react(path):
    if path.startswith('api/'):
        return jsonify(error='Not found'), 404
    index = os.path.join(REACT_DIR, 'index.html')
    if not os.path.exists(index):
        abort(503, description='React build missing. Run "npm run build" in /frontend.')
    return send_from_directory(REACT_DIR, 'index.html')

# =========================================================
# RUN APP
# =========================================================
if __name__ == '__main__':
    app.run(debug=True)
